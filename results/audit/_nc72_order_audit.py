# nc72 audit v2: first-citation order of main figs / supp figs / main tables / supp tables / references
# Source: fresh MS DOCX (results/CKI_Manuscript_NC.docx), body truncated before the References section.
# v2 fixes: comma-group capture (Tables 12, 13 / Figs. 15a, 16a); references via superscript runs.
import re, json
from docx import Document

DOCX = r"results/CKI_Manuscript_NC.docx"
doc = Document(DOCX)
paras = [p.text for p in doc.paragraphs]

ref_idx = None
for i, t in enumerate(paras):
    if t.strip().lower() in ("references", "reference"):
        ref_idx = i
        break
body = paras[:ref_idx]
print(f"total paras={len(paras)}, References starts at para {ref_idx}, body paras={len(body)}")
text = "\n".join(body)

NUM = r"[0-9]+[a-c]?"           # single number with optional block letter
RANGE = NUM + r"(?:\s*[\u2013-]\s*" + NUM + r")?"
GROUP = RANGE + r"(?:\s*,\s*" + RANGE + r")*"

def expand_group(s):
    out = []
    for tok in re.split(r",\s*", s):
        tok = tok.strip()
        if not tok:
            continue
        m = re.match(r"^(\d+)[a-z]?\s*[\u2013-]\s*(\d+)[a-z]?$", tok)
        if m:
            out.extend(range(int(m.group(1)), int(m.group(2)) + 1))
        else:
            m2 = re.match(r"^(\d+)", tok)
            if m2:
                out.append(int(m2.group(1)))
    return out

def first_order(pattern, flags=0):
    hits = []
    seq = 0
    for m in re.finditer(pattern, text, flags):
        for n in expand_group(m.group(1)):
            hits.append((m.start(), seq, n))
            seq += 1
    hits.sort(key=lambda x: (x[0], x[1]))
    seen, order = set(), []
    for pos, s, n in hits:
        if n not in seen:
            seen.add(n)
            order.append((n, pos))
    return order

pats = {
    "main_fig":   r"(?<!Supplementary )(?<!\w)Figs?\.?\s+(" + GROUP + r")",
    "supp_fig":   r"Supplementary Figs?\.?\s+(" + GROUP + r")",
    "main_table": r"(?<!Supplementary )(?<!\w)Tables?\s+(" + GROUP + r")",
    "supp_table": r"Supplementary Tables?\s+(" + GROUP + r")",
}

report = {}
for name, pat in pats.items():
    order = first_order(pat)
    seq = [n for n, _ in order]
    mono = seq == list(range(1, len(seq) + 1))
    report[name] = {"first_appearance_order": seq, "n_distinct": len(seq),
                    "strict_1..N_monotone": mono,
                    "positions": {str(n): pos for n, pos in order}}
    print(f"\n== {name} == distinct={len(seq)} monotone={mono}")
    print("  order:", seq)
    if name in ("supp_fig", "supp_table"):
        for n, pos in order:
            ctx = text[max(0, pos - 60):pos + 55].replace("\n", " | ")
            print(f"    #{n} @ {pos}: ...{ctx}...")

# --- references: superscript runs in body paragraphs ---
ref_hits = []   # (para_idx, run_order, text)
for pi, p in enumerate(doc.paragraphs[:ref_idx]):
    for ri, r in enumerate(p.runs):
        if r.font.superscript and r.text.strip():
            ref_hits.append((pi, ri, r.text.strip()))
ref_order, seen = [], set()
ref_flat = []
for pi, ri, t in ref_hits:
    for n in expand_group(t.replace("\u2013", "-")):
        ref_flat.append((pi, n, t))
        if n not in seen:
            seen.add(n)
            ref_order.append(n)
mono = ref_order == list(range(1, len(ref_order) + 1))
report["reference"] = {"first_appearance_order": ref_order, "n_distinct": len(ref_order),
                       "strict_1..N_monotone": mono, "n_sup_runs": len(ref_hits)}
print(f"\n== reference == superscript runs={len(ref_hits)} distinct={len(ref_order)} monotone={mono}")
print("  order:", ref_order)
print("  raw superscript runs:", [t for _, _, t in ref_hits][:80])

# bracket check: any literal [N] groups in body (should be none / non-citation only)
br = re.findall(r"\[([0-9][0-9,.\s\u2013+-]*)\]", text)
print("\n  bracketed numeric groups in body (sanity):", br[:20])

with open(r"results/audit/_nc72_order_audit.json", "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=2)
print("\nwrote results/audit/_nc72_order_audit.json")
