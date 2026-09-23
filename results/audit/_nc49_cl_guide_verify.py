"""Verify regenerated Cover Letter and Reproducibility Guide (v49 fix round)."""
from docx import Document

checks = []

def chk(name, cond, detail=''):
    checks.append((name, 'PASS' if cond else 'FAIL', detail))

# ---------- Cover Letter ----------
cl = Document('results/CKI_NC_Cover_Letter.docx')
clfull = '\n'.join(p.text for p in cl.paragraphs)
wc = len(clfull.split())

chk('CL word count <= 530', wc <= 530, f'{wc} words')
chk('CL no rebuttal framing', 'Previously raised concerns' not in clfull)
chk('CL no 5th of 5', '5th of 5' not in clfull)
chk('CL new-evidence paragraph', 'Two properties of this work are worth stating explicitly' in clfull)
chk('CL change-detection framing (no classification AUC)',
    'specificity-first index' in clfull and '0.680' not in clfull
    and 'dynamic cell-state changes' in clfull)
chk('CL drift claim qualified',
    'lowest false-report rate among the continuous divergence metrics' in clfull)
chk('CL no "least of seven metrics"', 'least of seven metrics' not in clfull)
chk('CL below raw JS 10/10', 'below raw JS in all ten cell classes' in clfull)
chk('CL 3,567', '3,567' in clfull)
chk('CL no 3,563/3,596', '3,563' not in clfull and '3,596' not in clfull)
chk('CL CC new ranges',
    '1.10\u20132.46' in clfull and '1.3\u20133.3-fold' in clfull
    and '1.13\u20132.46' not in clfull and '2.1\u20133.6-fold' not in clfull)
chk('CL no baseline-driven/baseline-associated',
    'baseline-driven' not in clfull and 'baseline-associated' not in clfull)
chk('CL KRAS dual-adjusted upgrade',
    'after purity and smoking adjustment' in clfull)
chk('CL EGFR dissolved',
    'whose apparent association dissolved under purity adjustment' in clfull)
chk('CL Fig. 3 / Fig. 4 refs', '(Fig. 3)' in clfull and '(Fig. 4)' in clfull)

# ---------- Reproducibility Guide ----------
gd = Document('results/CKI_Reproducibility_Guide_NC.docx')
gdfull = '\n'.join(p.text for p in gd.paragraphs)

chk('Guide drift section = Result 4 (Fig. 3)',
    'Real-Data Drift Calibration' in gdfull and 'Result 4 (Fig. 3)' in gdfull)
chk('Guide TCGA = Result 5 (Fig. 4)', 'Result 5 (Fig. 4)' in gdfull)
chk('Guide cross-organ = Result 6, Fig. 5', 'Cross-organ conservation (Result 6, Fig. 5)' in gdfull)
chk('Guide brain = Result 7 (Fig. 6)', 'Result 7 (Fig. 6)' in gdfull)
chk('Guide 3,567', '3,567 samples' in gdfull)
chk('Guide no 3,563; 3,596 only in attrition breakdown (v49.14)',
    '3,563' not in gdfull and gdfull.count('3,596') == 1
    and 'of the 3,596 expression-matrix samples' in gdfull)
chk('Guide CC attrition breakdown (v49.14)',
    '3 do not appear in the assembled pair table' in gdfull
    and '3,593 unique barcodes' in gdfull
    and '29 expression-matrix samples excluded' not in gdfull)
chk('Guide mean-ratio caliber',
    'mean(omega_NN) / mean(omega_TT)' in gdfull and 'Fig. 4a' in gdfull)
chk('Guide no median NN/TT caliber',
    'median(omega_NN) / median(omega_TT)' not in gdfull)
chk('Guide SF12 -> SF11 (2 places)',
    'Supplementary Fig. 12' not in gdfull
    and gdfull.count('Supplementary Fig. 11') >= 1)
chk('Guide SF10 -> SF3', '(Supplementary Fig. 3)' in gdfull
    and '80_kang_demo_figure.py (Supplementary Fig. 10)' not in gdfull)
chk('Guide 5.10 v49 Analyses', '5.10 v49 Analyses' in gdfull)
chk('Guide nc49 scripts listed',
    'nc49_pilot_kang_techrep.py' in gdfull and 'nc49_brain_drift_ladder.py' in gdfull
    and 'nc49_tcga_main.py' in gdfull and 'nc49_pilot_lihc_cox.py' in gdfull
    and 'nc49_fig_drift_ladder.py' in gdfull and 'nc49_fig_tcga.py' in gdfull)
chk('Guide nc49 outputs listed',
    'results/nc49_pilot_kang_techrep.csv' in gdfull
    and 'results/nc49_brain_drift_ladder.csv' in gdfull
    and 'results/nc49_tcga_pancancer.csv' in gdfull
    and 'results/nc49_tcga_luad_mutation.csv' in gdfull
    and 'results/nc49_pilot_lihc_cox.csv' in gdfull)
chk('Guide v49 checklist entries',
    'Verify drift-calibration spot values' in gdfull
    and 'Verify TCGA sample count n = 3,567' in gdfull
    and 'NN/TT mean ratios 1.10\u20132.46' in gdfull
    and 'v49: Verify' not in gdfull)
chk('Guide NEW-5 brain ladder B values',
    gdfull.count('B = 100 for T1, 30 for T2 and T3') == 2
    and 'n-matched nulls (B = 200)' not in gdfull
    and 'cell-shuffle null (B = 200); calibration ratio' not in gdfull)
chk('Guide NEW-4 iteration labels cleared',
    'v44 Blind-Review' not in gdfull
    and '5.8 Blind-Review Analyses' in gdfull)
chk('Guide baseline-/denominator-driven zero',
    'baseline-driven' not in gdfull and 'denominator-driven' not in gdfull)
chk('Guide EGFR admixture sync',
    'an admixture artefact, Section 5.10c' in gdfull)
chk('Guide SI param-justification ref updated',
    'Section 3.9 of the Supplementary Information (Parameter Justification)' in gdfull
    and 'Section 3.11 of the Supplementary Information (Parameter Justification)' not in gdfull)
chk('Guide Jaccard honest framing',
    'marker Jaccard is lower still on the FPR statistic (19.9%)' in gdfull)
chk('Guide B-group scripts listed',
    'nc49_tcga_purity.py' in gdfull and 'nc49_tcga_luad_smoking.py' in gdfull)
chk('Guide B-group outputs listed',
    'results/nc49_tcga_purity.csv' in gdfull
    and 'results/nc49_tcga_admix_scores.csv' in gdfull
    and 'results/nc49_tcga_luad_smoking.csv' in gdfull)
chk('Guide kf_composition listed',
    'nc49_tcga_kf_composition.py' in gdfull
    and 'results/nc49_tcga_kf_composition.csv' in gdfull)

print('===== CL + Guide checks =====')
nfail = 0
for name, status, detail in checks:
    print(f'[{status}] {name} {detail}')
    if status == 'FAIL':
        nfail += 1
print(f'TOTAL: {len(checks)} checks, {nfail} failures')
raise SystemExit(1 if nfail else 0)
