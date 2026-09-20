# -*- coding: utf-8 -*-
"""验证生成的 CKI_Manuscript_NC.docx：新节存在、图注编号 1-7 顺序、关键数字出现。"""
from docx import Document

doc = Document(r"C:/Users/KnightZ/Desktop/细胞受选择/results/CKI_Manuscript_NC.docx")
paras = [p.text for p in doc.paragraphs]
out = []

# 1. new section heading
hits = [i for i, t in enumerate(paras)
        if "neutral-drift calibration" in t.lower()]
out.append(f"'neutral-drift calibration' paragraphs at: {hits}")
for i in hits:
    out.append(f"  [{i}] {paras[i][:100]}")

# 2. figure legends order
figs = [(i, t[:60]) for i, t in enumerate(paras)
        if t.startswith("Figure ") and t[7:8].isdigit()]
out.append("\nFigure legends in order:")
for i, t in figs:
    out.append(f"  [{i}] {t}")

# 3. main figure references in body (Fig. N)
import re
refs = {}
for i, t in enumerate(paras):
    for m in re.finditer(r"Fig\. (\d)[a-z]?(?![\d])", t):
        if "Supplementary" not in t[max(0, m.start()-20):m.start()]:
            refs.setdefault(m.group(1), []).append(i)
out.append("\nMain Fig. references by number (excluding Supplementary):")
for k in sorted(refs):
    out.append(f"  Fig. {k}: paragraphs {refs[k]}")

# 4. key numbers present
checks = ["0.963", "28.6%", "45.2%", "44.1%", "40.6%", "19.9%", "90.9%",
          "1.04", "1.80", "2.98", "36.7%", "23.3%", "2,161", "1,089",
          "1,656", "137-fold", "21.9", "9.7", "TODO-nc49-drift",
          "n-matched"]
txt = "\n".join(paras)
out.append("\nKey-number presence:")
for c in checks:
    out.append(f"  {c!r}: {txt.count(c)} occurrence(s)")

# 5. results heading sequence (level-2 headings in Results)
out.append("\nAll level-2 style headings around Results:")
for i, p_ in enumerate(doc.paragraphs):
    if p_.style.name.startswith("Heading 2"):
        out.append(f"  [{i}] {p_.text[:80]}")

open(r"C:/Users/KnightZ/Desktop/细胞受选择/results/audit/_nc49_docx_verify.txt",
     "w", encoding="utf-8").write("\n".join(out))
print("done")
