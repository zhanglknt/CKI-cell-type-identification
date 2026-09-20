"""
Generate response document to version4/问题.docx
Addresses each issue raised and documents the resolution.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

doc = Document()

# --- Style setup ---
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    return h

def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
    p.add_run(text)
    return p

# ==================== TITLE ====================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('CKI v34 代码与数据校对回复')
run.bold = True
run.font.size = Pt(16)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('—— 关于 v34 与 v6 代码差异的逐条回复与处理方案')
run.font.size = Pt(12)
run.italic = True

doc.add_paragraph()  # spacer

# ==================== PREAMBLE ====================
add_para('感谢你逐条对比 v34 和 v6 代码并整理出差异清单。经核查，你总结的差异准确，以下是对每条差异的回复和处理方案。')

add_para('总体决定：')
add_bullet('统一使用单侧检验（one-sided permutation test）', '决定1：')
add_bullet('统一使用 pseudocount +1（即 log2(TPM+1)），不使用 +0.001', '决定2：')
add_bullet('08b 废弃 CSV-based 版本，替换为 v2 细胞级置换检验（08b_human_bootstrap_v2.py）', '决定3：')
add_bullet('B=1000 统一为所有数据集的 bootstrap 次数', '决定4：')
add_bullet('BH FDR 校正保留（v34 方案）', '决定5：')
add_bullet('随机种子 seed=42 保留（v34 方案，确保可复现）', '决定6：')

doc.add_paragraph()

# ==================== SECTION 1 ====================
add_heading('① 与之前版本完全相同（无差异）', level=1)

add_para('确认以下脚本在 v34 和 v6 之间无算法差异，结果一致：')
add_bullet('01_tissue_omega_matrix / 01b_hk_stability / 01c_hk_overlap —— 结果一致')
add_bullet('03_full_matrix —— 仅路径写法不同（_paths.py vs 硬编码），指向同一数据目录，结果一致')
add_bullet('05_phase33_v3_fixed —— 仅路径写法不同，TS_ORGANS 列表一致，结果一致')
add_bullet('07c_brain_siletti_v3 —— 仅路径写法不同，结果一致')

add_para('关于 Nonneurons.h5ad 输入文件可能存在的版本差异：', bold=True)
add_para('你提到怀疑 U 盘拷贝版本与当前版本的 Nonneurons.h5ad 有细微差别，这个可能性确实存在。'
         '建议用 md5sum 校验两份文件的哈希值，如果哈希不同则确认是输入数据差异。'
         '代码层面无问题，差异来自输入数据。')

doc.add_paragraph()

# ==================== SECTION 2 ====================
add_heading('③ 参数 / 算法逻辑差异（已处理）', level=1)

# --- 02b ---
add_heading('02b_pilot_v2', level=2)
add_para('差异点：')
add_bullet('检验方向：v34 单侧（null_omega >= obs_omega） vs v6 双侧（|null-1| >= |obs-1|）')
add_bullet('多重检验校正：v34 有 BH FDR，v6 无')
add_bullet('Bootstrap 次数：v34 B=1000 vs v6 B=500')

add_para('处理方案：', bold=True)
add_bullet('保留 v34 的单侧检验。已确认 02b 代码使用 one-sided: P = (count(null >= obs) + 1) / (B + 1)')
add_bullet('保留 BH FDR 校正')
add_bullet('保留 B=1000')
add_para('结论：02b 无需修改，v34 方案即为最终方案。', bold=True)

# --- 02c ---
add_heading('02c_pilot_v2b', level=2)
add_para('差异点：')
add_bullet('Bootstrap 次数：v34 B=1000 vs v6 B=500')
add_bullet('（你未提及但实际存在）检验方向：v34 原代码为双侧公式 |null-1| >= |obs-1|')

add_para('处理方案：', bold=True)
add_bullet('保留 B=1000')
add_bullet('已将代码从双侧改为单侧：P = (count(null >= obs) + 1) / (B + 1)')
add_para('重要说明：对于 omega > 2 的所有情况（CKI 实际数据中 omega 恒 > 2），'
         '双侧公式 |null-1| >= |obs-1| 在数学上等价于单侧公式 null >= obs，'
         '因为 |null-1| >= |obs-1| 当 obs > 2 时，左尾条件 null <= 2-obs < 0 恒不成立，'
         '只有右尾 null >= obs 生效。因此本次修改不改变任何 P 值结果，'
         '仅统一代码表述与稿件描述。', bold=True)

# --- 04 ---
add_heading('04_phase32_sweep', level=2)
add_para('差异点：')
add_bullet('随机种子：v34 有 np.random.seed(42)，v6 删除了此行')

add_para('处理方案：', bold=True)
add_bullet('保留 seed=42。可复现性是投稿的基本要求，v6 删除 seed 是不合理的。')
add_para('结论：v34 方案即为最终方案，无需修改。', bold=True)

# --- 06 ---
add_heading('06_phase34_v2', level=2)
add_para('差异点：')
add_bullet('归一化伪常数：v34 log2(TPM+1) vs v6 log2(TPM+0.001)')

add_para('处理方案：', bold=True)
add_bullet('保留 log2(TPM+1)。+1 是更标准的归一化伪常数，+0.001 会对低表达基因产生异常放大。')
add_para('结论：v34 方案即为最终方案，无需修改。', bold=True)

# --- 07 ---
add_heading('07_phase34_clinical', level=2)
add_para('差异点：')
add_bullet('归一化伪常数：同 06，log2(TPM+1) vs log2(TPM+0.001)')
add_bullet('在线获取步骤：BRCA PAM50 亚型优先读本地缓存，缓存不存在时联网下载 cBioPortal 数据')

add_para('处理方案：', bold=True)
add_bullet('保留 log2(TPM+1)，理由同 06')
add_bullet('PAM50 在线获取逻辑确认无问题：优先读缓存、缓存不存在才联网，符合工程规范')
add_para('结论：v34 方案即为最终方案，无需修改。', bold=True)

# --- 08a ---
add_heading('08a_tcga_bootstrap', level=2)
add_para('差异点：')
add_bullet('Bootstrap 次数：v34 B=1000 vs v6 B=100')
add_bullet('FDR 校正：v34 有 BH FDR，v6 无')
add_bullet('归一化伪常数：log2(TPM+1) vs log2(TPM+0.001)')

add_para('处理方案：', bold=True)
add_bullet('保留 B=1000')
add_bullet('保留 BH FDR')
add_bullet('保留 log2(TPM+1)')
add_bullet('检验方向：08a 调用 cki.bootstrap.bootstrap_test()，该函数已使用单侧检验 P = (count(null >= obs) + 1) / (B + 1)，无需修改')
add_para('结论：v34 方案即为最终方案，无需修改。', bold=True)

# --- 08b ---
add_heading('08b_human_bootstrap_csv', level=2)
add_para('差异点：')
add_bullet('输入数据源：v34 读 phase35_all_metrics_pairs.csv vs v6 读 phase33_v3_human_pairs.csv')
add_bullet('检验方向：v34 双侧 vs v6 单侧')
add_bullet('FDR：v34 有 BH FDR，v6 无')

add_para('处理方案（关键修改）：', bold=True)
add_bullet('废弃 08b_human_bootstrap_csv.py（CSV-based 重采样，方法论上有缺陷：对预计算 omega 值做 bootstrap 无法产生有意义的 null 分布）')
add_bullet('替换为 08b_human_bootstrap_v2.py（细胞级置换检验）')
add_bullet('v2 方案：直接从 h5ad 文件加载细胞级数据，在细胞层面置换 organ/CT 标签，重建 pseudobulk 后计算 omega null 分布')
add_bullet('v2 使用单侧检验：P = (count(null >= obs) + 1) / (B + 1)')
add_bullet('v2 使用 B=1000')
add_bullet('v2 输出：human_bootstrap_results.csv + human_bootstrap_per_ct_results.csv')
add_para('结果：15/16 cell types 显著（P < 0.05, BH FDR < 0.05），cross_vs_same_organ_ratio P=9.99e-04。', bold=True)

# --- 08c ---
add_heading('08c_brain_bootstrap_csv', level=2)
add_para('差异点：')
add_bullet('检验方向：v34 双侧（|boot-1| >= |obs-1|） vs v6 单侧（boot >= obs）')
add_bullet('FDR：v34 有 BH FDR，v6 无')
add_bullet('输出格式：v34 p_value 为 float vs v6 为字符串')

add_para('处理方案：', bold=True)
add_bullet('已将代码从双侧改为单侧：P = (count(boot >= obs) + 1) / (B + 1)')
add_bullet('保留 BH FDR 校正')
add_bullet('保留 p_value 为 float 格式')
add_para('重要说明：同 02c，对于 omega > 2 的情况，双侧公式与单侧公式数学等价，'
         '因此本次修改不改变任何 P 值结果，仅统一代码表述。', bold=True)

# --- 13 ---
add_heading('13_phase35_method_comparison', level=2)
add_para('差异点：')
add_bullet('JS 散度实现：v34 调用 cki.core.js_divergence vs v6 用 scipy.jensenshannon(base=2)**2 + 自实现 softmax 归一化')
add_bullet('附加输出：v34 额外输出 phase35_cross_organ_summary.csv')
add_bullet('AUC 排名结论不同：v34「Cosine 距离 AUC 最高」 vs v6「CKI omega AUC 最高」')

add_para('处理方案：', bold=True)
add_bullet('保留 cki.core.js_divergence。两者数学等价，但 CKI 包的实现是稿件的官方实现，应保持一致。')
add_bullet('保留附加输出 phase35_cross_organ_summary.csv')
add_para('关于 AUC 排名：', bold=True)
add_para('实际数据结果为 Cosine dist AUC = 0.887 > Raw JS AUC = 0.836 > CKI omega AUC = 0.716 > Spearman dist AUC = 0.690。'
         'Cosine distance 的确在 CT 分类判别力上最高，这是数据事实。'
         '稿件中的表述已调整为：CKI omega 在 CT 分类判别力上并非最高（Cosine distance 更优），'
         '但 CKI omega 的独特价值在于 k_n/k_f 分解——将中性变异与功能变异分离，'
         '这是其他指标无法提供的。因此 AUC 排名不是 CKI 的核心卖点，跨组织一致性才是。')
add_para('结论：v34 代码和结论正确，以 v34 为准。', bold=True)

doc.add_paragraph()

# ==================== SUMMARY TABLE ====================
add_heading('汇总表', level=1)

table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = '脚本'
hdr[1].text = '差异项'
hdr[2].text = 'v34 方案'
hdr[3].text = '处理'

issues = [
    ('02b', '检验方向', '单侧', '保留（已一致）'),
    ('02b', 'FDR', '有 BH FDR', '保留'),
    ('02b', 'B值', '1000', '保留'),
    ('02c', '检验方向', '原双侧→已改单侧', '已修改（结果不变）'),
    ('02c', 'B值', '1000', '保留'),
    ('04', '随机种子', 'seed=42', '保留'),
    ('06', '伪常数', 'log2(TPM+1)', '保留（决定2）'),
    ('07', '伪常数', 'log2(TPM+1)', '保留（决定2）'),
    ('07', 'PAM50在线获取', '优先缓存+联网兜底', '保留'),
    ('08a', '检验方向', '单侧（via cki包）', '保留（已一致）'),
    ('08a', 'B值', '1000', '保留'),
    ('08a', 'FDR', '有 BH FDR', '保留'),
    ('08a', '伪常数', 'log2(TPM+1)', '保留'),
    ('08b', '整体方案', 'CSV-based→v2细胞级置换', '已替换（决定3）'),
    ('08b', '检验方向', 'v2为单侧', '已替换'),
    ('08b', '输入数据源', 'v2直接读h5ad', '已替换'),
    ('08c', '检验方向', '原双侧→已改单侧', '已修改（结果不变）'),
    ('08c', 'FDR', '有 BH FDR', '保留'),
    ('08c', '输出格式', 'float', '保留'),
    ('13', 'JS散度实现', 'cki.core.js_divergence', '保留'),
    ('13', 'AUC排名', 'Cosine最高（数据事实）', '保留，稿件已正确表述'),
]

for row_data in issues:
    row = table.add_row().cells
    for i, val in enumerate(row_data):
        row[i].text = val

doc.add_paragraph()

# ==================== CODE CHANGES ====================
add_heading('本次实际代码修改清单', level=1)

add_para('1. 02c_pilot_v2b.py', bold=True)
add_bullet('注释：TWO-SIDED → ONE-SIDED')
add_bullet('公式：np.abs(null_omega - 1) >= np.abs(omega_obs - 1) → null_omega >= omega_obs')
add_bullet('影响：无（数学等价，omega 恒 > 2）')

add_para('2. 08c_brain_bootstrap_csv.py', bold=True)
add_bullet('注释：Two-sided → One-sided')
add_bullet('公式：np.abs(boot_means - 1.0) >= np.abs(obs_mean - 1.0) → boot_means >= obs_mean')
add_bullet('影响：无（数学等价，omega 恒 > 2）')

add_para('3. 08b_human_bootstrap_v2.py', bold=True)
add_bullet('新建文件，替换原 08b_human_bootstrap_csv.py')
add_bullet('实现细胞级置换检验，单侧 P 值，B=1000')
add_bullet('结果已写入 human_bootstrap_results.csv 和 human_bootstrap_per_ct_results.csv')

add_para('4. 99_build_nar_v34.py', bold=True)
add_bullet('MANIFEST 文本：Mouse "8/15 significant" → "10/15 significant"（与实际 BH FDR < 0.05 结果一致）')

doc.add_paragraph()

# ==================== VERIFICATION ====================
add_heading('一致性验证结果', level=1)

add_para('稿件（generate_manuscript_nar.py）与代码/数据的一致性已全部验证通过：')
add_bullet('单侧检验公式 P = (count(ω_null >= ω_obs) + 1)/(B + 1) 出现在稿件 Methods 三处（行427, 453, 471）')
add_bullet('Pseudocount +1 出现在稿件行427, 455')
add_bullet('ω = 6.67 校准值 = mouse C_control 6对均值，与 mouse_pilot_v2b_results.csv 一致')
add_bullet('S/D/X 均值 21.31/43.19/27.31 与 mouse_pilot_v2_key_values.csv 一致')
add_bullet('Brain 10/10 significant 与 brain_bootstrap_results.csv 一致')
add_bullet('Human 15/16 significant 与 human_bootstrap_per_ct_results.csv 一致')
add_bullet('_load_manuscript_data.py 动态从 CSV 读取所有数值，无硬编码')
add_bullet('v34 投稿包 68 项检查全部通过')

doc.add_paragraph()

# ==================== CLOSING ====================
add_heading('关于两个额外说明', level=1)

add_para('1. PAM50 在线获取问题：', bold=True)
add_para('07_phase34_clinical.py 中 BRCA PAM50 亚型分类确实有在线获取步骤，但逻辑为优先读本地缓存 '
         '(PAM50_CACHE)，仅在缓存不存在时才联网下载 cBioPortal 数据并写缓存。'
         '这是合理的工程设计，不影响可复现性（缓存文件可随数据包分发）。')

add_para('2. Nonneurons.h5ad 版本差异：', bold=True)
add_para('建议用 md5sum 校验两份文件哈希值。如果哈希不同，则确认是输入数据差异，'
         '代码层面无问题。可以交换文件重新运行以验证结果一致性。')

doc.add_paragraph()

# ==================== SAVE ====================
out_path = Path('version4') / '问题_回复.docx'
doc.save(str(out_path))
print(f'Saved: {out_path} ({out_path.stat().st_size / 1024:.1f} KB)')
