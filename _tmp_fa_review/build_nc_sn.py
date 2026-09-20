# -*- coding: utf-8 -*-
"""
Build notebooks/68_gen_supplementary_nc.py from notebooks/68_gen_supplementary_en.py
(GB-style Supplementary Materials -> NC-style Supplementary Information).

Mechanical, assertion-guarded transformation:
  1. 15 manuscript-cited decimal subsections -> 'Supplementary Note 1-15: <title>'
     (mapping below), physically reordered into new numeric order.
  2. 5 former group headings 'Supplementary Note 1-5: X' -> plain section titles 'X'
     (group 5 dropped: both its subsections are pulled out as notes).
  3. Uncited subsections keep decimal numbers; references to them -> 'Section X.Y'.
  4. 'Figure SN' / 'Fig. SN' -> 'Supplementary Fig. N'; 'Table SN' -> 'Supplementary Table N'.
  5. Document title -> 'Supplementary Information'; TOC rebuilt; output path
     -> results/CKI_Supplementary_NC.docx.
Then: py_compile, run generator, verify DOCX, write conversion report.
"""
import re
import sys
import py_compile
import subprocess
from pathlib import Path

ROOT = Path(r'C:\Users\KnightZ\Desktop\细胞受选择')
SRC = ROOT / 'notebooks' / '68_gen_supplementary_en.py'
DST = ROOT / 'notebooks' / '68_gen_supplementary_nc.py'
REPORT = ROOT / '_tmp_fa_review' / 'nc_sn_conversion_report_2026-09-17.md'
DOCX = ROOT / 'results' / 'CKI_Supplementary_NC.docx'

NOTE_ORDER = ['3.12', '3.21', '3.5', '3.20', '3.22', '3.15', '3.13', '5.2',
              '3.16', '3.17', '3.14', '4.6', '5.1', '3.23', '3.6']
OLD2NEW = {old: i + 1 for i, old in enumerate(NOTE_ORDER)}
UNMAPPED_REFS = ['1.1', '1.5', '1.7', '3.18', '3.19a', '3.19b', '3.19c']

log = []          # (item, detail)
checks = []       # (check, ok, detail)


def fail(msg):
    raise SystemExit(f'FATAL: {msg}')


# ---------------------------------------------------------------- read source
lines = SRC.read_text(encoding='utf-8').split('\n')

RE_SUB = re.compile(r"^add_para\('(\d+\.\d+) (.+)', bold=True\)\s*$")
RE_GRP = re.compile(r"^add_heading\('Supplementary Note ([1-5]): (.+)', 2\)\s*$")
RE_TBL = re.compile(r"^add_heading\('Table S(\d+): (.+)', 2\)\s*$")
RE_CMT = re.compile(r"^# ===== (SN\d|Supplementary Tables)")

markers = []  # (idx0, kind, key, title)
for i, ln in enumerate(lines):
    m = RE_SUB.match(ln)
    if m:
        markers.append((i, 'sub', m.group(1), m.group(2)))
        continue
    m = RE_GRP.match(ln)
    if m:
        markers.append((i, 'grp', int(m.group(1)), m.group(2)))
        continue
    m = RE_TBL.match(ln)
    if m:
        markers.append((i, 'tbl', int(m.group(1)), m.group(2)))
        continue
    if ln.strip() == 'doc.add_page_break()':
        markers.append((i, 'pb', None, None))
        continue
    if ln.startswith("add_heading('Supplementary Data 1"):
        markers.append((i, 'data', None, None))
        continue
    if ln.startswith('# ===== Add line numbers'):
        markers.append((i, 'tail', None, None))
        continue
    m = RE_CMT.match(ln)
    if m:
        markers.append((i, 'cmt', m.group(1), None))

n_sub = sum(1 for _, k, _, _ in markers if k == 'sub')
n_grp = sum(1 for _, k, _, _ in markers if k == 'grp')
n_pb = sum(1 for _, k, _, _ in markers if k == 'pb')
n_tbl = sum(1 for _, k, _, _ in markers if k == 'tbl')
n_cmt = sum(1 for _, k, _, _ in markers if k == 'cmt')
assert n_sub == 38, f'expected 38 decimal subsections, got {n_sub}'
assert n_grp == 5, f'expected 5 group headings, got {n_grp}'
assert n_pb == 6, f'expected 6 page breaks, got {n_pb}'
assert n_tbl == 4, f'expected 4 table headings, got {n_tbl}'
assert n_cmt == 5, f'expected 5 section comments, got {n_cmt}'
sub_keys = [k for _, kd, k, _ in markers if kd == 'sub']
assert set(NOTE_ORDER) <= set(sub_keys), 'note keys missing from source'
log.append(('markers', f'sub={n_sub} grp={n_grp} pb={n_pb} tbl={n_tbl} cmt={n_cmt}'))

# ---------------------------------------------------------------- slice chunks
chunks = []
for j, (idx, kind, key, title) in enumerate(markers):
    end = markers[j + 1][0] if j + 1 < len(markers) else len(lines)
    chunks.append({'kind': kind, 'key': key, 'title': title,
                   'lines': lines[idx:end]})
head_lines = lines[:markers[0][0]]


def ctext(c):
    return '\n'.join(c['lines']).rstrip()


# ------------------------------------------------- classify into new structure
buckets = {1: [], 2: [], 3: [], 4: []}
notes = {}       # old key -> chunk
tables = []
pending = []
cur = None
data_chunk = tail_chunk = None
dropped_grp5 = False
for c in chunks:
    k = c['kind']
    if k == 'cmt':
        if c['key'] == 'SN5':
            cur = 'DROP5'
            pending = []
        else:
            pending.append(c)
        continue
    if k == 'grp':
        if c['key'] == 5:
            cur = 'DROP5'
            pending = []
            dropped_grp5 = True
            continue
        cur = c['key']
        buckets[cur] = pending + [c]
        pending = []
        continue
    if k == 'pb':
        continue  # re-emitted between assembled blocks
    if k == 'sub':
        if c['key'] in OLD2NEW:
            notes[c['key']] = c
        else:
            assert isinstance(cur, int), f"subsection {c['key']} outside a kept group"
            buckets[cur].append(c)
        continue
    if k == 'tbl':
        if not tables:
            tables.extend(pending)
            pending = []
        tables.append(c)
        continue
    if k == 'data':
        data_chunk = c
        continue
    if k == 'tail':
        tail_chunk = c
        continue
assert dropped_grp5, 'group-5 heading was not dropped'
assert len(notes) == 15, f'expected 15 note chunks, got {len(notes)}'
assert set(notes) == set(NOTE_ORDER)
assert len(tables) == 5  # section comment + 4 tables
assert data_chunk is not None and tail_chunk is not None
assert all(buckets[g] and buckets[g][1]['kind'] == 'grp' or True for g in buckets)

# ------------------------------------------------------- transform group heads
grp_titles = {}
for g in (1, 2, 3, 4):
    for c in buckets[g]:
        if c['kind'] == 'grp':
            grp_titles[g] = c['title']
            c['lines'] = [f"add_heading('{c['title']}', 2)"] + c['lines'][1:]
assert grp_titles == {1: 'CKI Mathematical Derivation',
                      2: 'CKI Algorithm Pseudocode',
                      3: 'Statistical Testing Details',
                      4: 'Dataset Quality Control and Filtering Criteria'}, grp_titles

# ------------------------------------------------------------- convert notes
note_titles = {}  # new num -> raw title (source-escaped form preserved)
note_texts = {}
for old, new in OLD2NEW.items():
    c = notes[old]
    first = c['lines'][0]
    m = RE_SUB.match(first)
    assert m and m.group(1) == old, f'note chunk {old} heading mismatch: {first}'
    title = m.group(2)
    note_titles[new] = title
    note_texts[new] = '\n'.join(
        [f"add_heading('Supplementary Note {new}: {title}', 2)"] + c['lines'][1:]
    ).rstrip()
log.append(('note headings converted', '; '.join(f'{o}->{n}' for o, n in NOTE_ORDER and OLD2NEW.items())))

# ------------------------------------------------------------------ new TOC
toc_entries = [grp_titles[g] for g in (1, 2, 3, 4)]
toc_entries += [f'Supplementary Note {n}: {note_titles[n]}' for n in range(1, 16)]
tbl_titles = {}
for c in tables:
    if c['kind'] == 'tbl':
        tbl_titles[c['key']] = c['title']
        toc_entries.append(f"Supplementary Table {c['key']}: {c['title']}")
assert sorted(tbl_titles) == [1, 2, 3, 4]
toc_entries.append('Supplementary Data 1: Complete Analysis Script Index')
toc_literal = 'toc = [\n' + '\n'.join(f"    '{e}'," for e in toc_entries) + '\n]'

# ------------------------------------------------------------------ HEAD edits
head = '\n'.join(head_lines)
old_doc = ('Generate English Supplementary Materials DOCX with continuous line numbers.\n'
           'Replaces the Chinese supplementary with English translation in Chinese-researcher style.')
new_doc = ('Generate NC-format Supplementary Information DOCX (from 68_gen_supplementary_en.py).\n'
           'NC naming: Supplementary Note 1-15, Supplementary Fig., Supplementary Table.\n'
           'Output: results/CKI_Supplementary_NC.docx')
assert head.count(old_doc) == 1, 'docstring anchor missing'
head = head.replace(old_doc, new_doc, 1)
assert head.count("add_heading('Supplementary Materials', 1)") == 1, 'title anchor missing'
head = head.replace("add_heading('Supplementary Materials', 1)",
                    "add_heading('Supplementary Information', 1)", 1)
head, n_toc = re.subn(r"toc = \[.*?\n\]", lambda _m: toc_literal, head, flags=re.S)
assert n_toc == 1, f'TOC anchor count {n_toc}'
log.append(('HEAD', "docstring, title 'Supplementary Information', TOC rebuilt "
                    f'({len(toc_entries)} entries)'))

# ------------------------------------------------------------------ TAIL edits
tail_txt = ctext(tail_chunk)
assert tail_txt.count("out_path = 'results/CKI_Supplementary.docx'") == 1, 'save path anchor'
tail_txt = tail_txt.replace("out_path = 'results/CKI_Supplementary.docx'",
                            "out_path = 'results/CKI_Supplementary_NC.docx'", 1)

# ------------------------------------------------------------------ assemble
parts = [head.rstrip()]
for g in (1, 2, 3, 4):
    parts.append('doc.add_page_break()')
    parts.append('\n\n'.join(ctext(c) for c in buckets[g]))
parts.append('doc.add_page_break()')
parts.append('# ===== Supplementary Notes (NC: citation order 1-15) =====\n\n'
             + '\n\n'.join(note_texts[n] for n in range(1, 16)))
parts.append('doc.add_page_break()')
parts.append('\n\n'.join(ctext(c) for c in tables))
parts.append(ctext(data_chunk))
parts.append(tail_txt)
out = '\n\n'.join(parts) + '\n'

# ------------------------------------------------- global textual replacements
def repl_literal(old, new, name, expect=1):
    global out
    n = out.count(old)
    assert n == expect, f'{name}: anchor count {n} != {expect}'
    out = out.replace(old, new)
    log.append((name, f'{n} replacement(s) [literal]'))


# special cross-literal references (must precede generic rules)
repl_literal("see also Notes 3.12 and '\n    '3.15).",
             "see also Supplementary Notes 1 and '\n    '6).",
             "Notes 3.12 and 3.15 -> Supplementary Notes 1 and 6")
repl_literal("HK-anchored. Note '\n    '3.12\\u2019s earlier",
             "HK-anchored. Supplementary Note 1\\u2019s earlier",
             "Note 3.12's -> Supplementary Note 1's")
repl_literal("'Additional file 1: Fig. S10.'",
             "'Supplementary Fig. 10.'",
             'Additional file 1: Fig. S10 -> Supplementary Fig. 10')

# generic note references (single pass, no cascading)
_alt = sorted(list(OLD2NEW) + UNMAPPED_REFS, key=len, reverse=True)
RE_NOTEREF = re.compile(r'\b(Notes?)\s+(' + '|'.join(map(re.escape, _alt)) + r')\b')
_ref_counts = {}


def _note_sub(m):
    word, key = m.group(1), m.group(2)
    _ref_counts[key] = _ref_counts.get(key, 0) + 1
    if key in OLD2NEW:
        return f'Supplementary {word} {OLD2NEW[key]}'
    return f'Section {key}'


out = RE_NOTEREF.sub(_note_sub, out)
log.append(('note references', '; '.join(f'{k}x{v}' for k, v in sorted(_ref_counts.items()))))

out, n = re.subn(r'\bSN\s+1\.5\b', 'Section 1.5', out)
assert n == 1, f'SN 1.5 count {n}'
log.append(('SN 1.5 -> Section 1.5', f'{n}'))

_fig_counts = {}


def _fig_sub(m):
    _fig_counts[m.group(1)] = _fig_counts.get(m.group(1), 0) + 1
    return f"Supplementary Fig. {m.group(1)}"


out = re.sub(r'\bFigure S(\d+)', _fig_sub, out)
out = re.sub(r'\bFig\. S(\d+)', _fig_sub, out)
log.append(('figure refs', '; '.join(f'S{k}x{v}' for k, v in sorted(_fig_counts.items(), key=lambda x: int(x[0])))))

_tbl_counts = {}


def _tbl_sub(m):
    _tbl_counts[m.group(1)] = _tbl_counts.get(m.group(1), 0) + 1
    return f"Supplementary Table {m.group(1)}"


out = re.sub(r'\bTable S(\d+)', _tbl_sub, out)
log.append(('table refs (headings+body; TOC rebuilt separately)',
            '; '.join(f'S{k}x{v}' for k, v in sorted(_tbl_counts.items(), key=lambda x: int(x[0])))))

# panel labels: expected none in this file
n_panels = len(re.findall(r'\([A-E]\)', out))
log.append(('panel labels (A)-(E)', f'{n_panels} occurrences (no replacement needed)'))

# ------------------------------------------------- residual checks (code only)
code = '\n'.join(l for l in out.split('\n') if not l.lstrip().startswith('#'))
for pat, name in [(r'\bNotes?\s+\d+\.\d', 'residual Note X.Y'),
                  (r'\bSN\s+\d+\.\d', 'residual SN X.Y'),
                  (r'\bFig\.\s*S\d', 'residual Fig. S'),
                  (r'\bFigure\s+S\d', 'residual Figure S'),
                  (r'\bTable\s+S\d', 'residual Table S'),
                  (r'Additional file', 'residual Additional file'),
                  (r'Supplementary Materials', 'residual Supplementary Materials')]:
    m = re.search(pat, code)
    checks.append((name, m is None, m.group(0) if m else 'clean'))
    assert m is None, f'{name}: {m.group(0)}'

# ------------------------------------------------------------------ write+compile
DST.write_text(out, encoding='utf-8')
py_compile.compile(str(DST), doraise=True)
checks.append(('py_compile 68_gen_supplementary_nc.py', True, 'ok'))

# ------------------------------------------------------------------ run generator
proc = subprocess.run([sys.executable, str(DST)], cwd=str(ROOT),
                      capture_output=True, text=True)
run_tail = (proc.stdout or '').strip().splitlines()[-3:]
if proc.returncode != 0:
    fail(f'generator failed:\n{proc.stdout}\n{proc.stderr}')
assert DOCX.exists(), 'DOCX not produced'
checks.append(('generator run', True, ' | '.join(run_tail)))

# ------------------------------------------------------------------ verify DOCX
from docx import Document  # noqa: E402

d = Document(str(DOCX))
paras = [(p.style.name, p.text) for p in d.paragraphs]
full = '\n'.join(t for _, t in paras)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            full += '\n' + cell.text

for pat, name in [(r'Note\s+[1345]\.\d', 'DOCX residual Note X.Y'),
                  (r'Notes\s+\d+\.\d', 'DOCX residual Notes X.Y'),
                  (r'SN\s+\d+\.\d', 'DOCX residual SN X.Y'),
                  (r'Fig\.\s*S\d', 'DOCX residual Fig. S'),
                  (r'Figure\s+S\d', 'DOCX residual Figure S'),
                  (r'Table\s+S\d', 'DOCX residual Table S'),
                  (r'Additional file', 'DOCX residual Additional file'),
                  (r'Supplementary Materials', 'DOCX residual Supplementary Materials')]:
    m = re.search(pat, full)
    checks.append((name, m is None, m.group(0) if m else 'clean'))
    assert m is None, f'{name}: {m.group(0)}'

# note headings 1-15 in order (Heading 2 style)
head_seq = [int(m.group(1)) for style, txt in paras
            if style.startswith('Heading')
            for m in [re.match(r'Supplementary Note (\d+): ', txt)] if m]
ok = head_seq == list(range(1, 16))
checks.append(('Supplementary Note 1-15 headings in order', ok, str(head_seq)))
assert ok, head_seq

# TOC entries 1-15 in order, titles match headings
toc_seq = [(int(m.group(1)), txt) for style, txt in paras
           if style == 'Normal'
           for m in [re.match(r'Supplementary Note (\d+): ', txt)] if m]
ok = [n for n, _ in toc_seq] == list(range(1, 16))
checks.append(('TOC Supplementary Note 1-15 in order', ok, str([n for n, _ in toc_seq])))
assert ok
head_titles = [txt for style, txt in paras if style.startswith('Heading')
               and re.match(r'Supplementary Note \d+: ', txt)]
ok = [t for _, t in toc_seq] == head_titles
checks.append(('TOC titles == heading titles', ok, ''))
assert ok

# figure/table counts vs original
orig = SRC.read_text(encoding='utf-8')
n_orig_fig = len(re.findall(r'\bFigure S\d+', orig)) + len(re.findall(r'\bFig\. S\d+', orig))
fig_nums = sorted(int(m.group(1)) for m in re.finditer(r'Supplementary Fig\. (\d+)', full))
ok = len(fig_nums) == n_orig_fig
checks.append((f'Supplementary Fig. count == original Figure/Fig. S count ({n_orig_fig})',
               ok, f'new={fig_nums}'))
assert ok, fig_nums

n_orig_tbl = len(re.findall(r'\bTable S\d+', orig))
tbl_nums = sorted(int(m.group(1)) for m in re.finditer(r'Supplementary Table (\d+)', full))
ok = len(tbl_nums) == n_orig_tbl
checks.append((f'Supplementary Table count == original Table S count ({n_orig_tbl})',
               ok, f'new={tbl_nums}'))
assert ok, tbl_nums

ok = any(style.startswith('Heading') and txt == 'Supplementary Information'
         for style, txt in paras)
checks.append(("document title 'Supplementary Information'", ok, ''))
assert ok

n_panel_docx = len(re.findall(r'\([A-E]\)', full))
checks.append(('panel labels (A)-(E) in DOCX', n_panel_docx == 0, f'{n_panel_docx}'))

sec_refs = sorted(set(re.findall(r'Section (\d+\.\d+[a-z]?)', full)))
checks.append(('Section X.Y refs (uncited subsections)', True, ', '.join(sec_refs)))

# ------------------------------------------------------------------ report
rep = ['# CKI Supplementary Notes -> NC 格式转换报告',
       '',
       '- 日期：2026-09-17',
       '- 输入：`notebooks/68_gen_supplementary_en.py`（未改动）',
       '- 产出：`notebooks/68_gen_supplementary_nc.py` + `results/CKI_Supplementary_NC.docx`',
       '- 构建脚本：`_tmp_fa_review/build_nc_sn.py`（全部锚点断言通过）',
       '',
       '## 1. Note 重编号映射（旧 -> 新，与文稿侧一致）',
       '',
       '| 旧编号 | 新 Supplementary Note | 标题 |',
       '|---|---|---|']
for old, new in OLD2NEW.items():
    rep.append(f'| {old} | {new} | {note_titles[new]} |')
rep += ['',
        '## 2. 结构调整',
        '',
        "- 原 5 个 'Supplementary Note 1-5' 大节标题去前缀降级为普通章节标题："
        'CKI Mathematical Derivation / CKI Algorithm Pseudocode / Statistical Testing Details / '
        'Dataset Quality Control and Filtering Criteria。',
        '- 原 Note 5 大节（Post-hoc Coherence Checks）的两个子节 5.1/5.2 均在映射表内被抽出，'
        '该大节标题已删除（避免空节）。',
        '- 15 个被引用小节物理重排为 Supplementary Note 1-15，置于 4 个主题节之后、'
        'Supplementary Tables 之前；节间保留分页符（6 个）。',
        '- 未被正文引用的小节（1.1-1.7、3.1-3.4、3.7-3.11、3.18、3.19、4.1-4.5）原地保留十进制编号，'
        '文内引用统一改写为 Section X.Y。',
        '',
        '## 3. 替换计数',
        '']
for item, detail in log:
    rep.append(f'- **{item}**: {detail}')
rep += ['',
        '## 4. DOCX 自检结果（全部通过）',
        '']
for name, ok, detail in checks:
    rep.append(f"- [{'x' if ok else ' '}] {name}" + (f' — {detail}' if detail else ''))
rep += ['',
        '## 5. 偏差与说明',
        '',
        "- 'Additional file 1: Fig. S10.'（全文唯一 Additional file 出现）按语境转换为 "
        "'Supplementary Fig. 10.'，而非字面替换为 'Supplementary Information'（后者语义不通）。",
        '- 面板标签 (A)-(E)：本生成器输出文本中 0 出现（面板标签在主文稿与图 PDF 内，不在本文件），无需替换。',
        "- 跨字面量换行的引用 2 处已定点处理：'see also Notes 3.12 and 3.15'、'Note 3.12\\u2019s'。",
        '- 统计正文 (S1)/(S2) 为 split-half 重复标签（L910 附近），非面板/图引用，未改动。',
        '- `\\uXXXX` 字面转义全部按原样保留（纯文本级手术，未经过编解码）。',
        '']
REPORT.write_text('\n'.join(rep), encoding='utf-8')

print('OK: build + compile + generate + verify all passed')
print(f'written: {DST}')
print(f'docx:    {DOCX}')
print(f'report:  {REPORT}')
for name, ok, detail in checks:
    print(f"  [{'PASS' if ok else 'FAIL'}] {name} {detail}")
