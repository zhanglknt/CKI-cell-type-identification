# -*- coding: utf-8 -*-
"""verify_nc_guide.py — python-docx 提取 NC 指南全文，自检残留与计数。"""
import re
import sys
from docx import Document

DOCX = "C:/Users/KnightZ/Desktop/细胞受选择/results/CKI_Reproducibility_Guide_NC.docx"

doc = Document(DOCX)
parts = [p.text for p in doc.paragraphs]
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            parts.append(cell.text)
for sec in doc.sections:
    for p in sec.header.paragraphs:
        parts.append(p.text)
    for p in sec.footer.paragraphs:
        parts.append(p.text)
text = "\n".join(parts)

residual_pats = [
    r"Note\s+3\.", r"Note\s+4\.", r"Note\s+5\.",
    r"Fig\.\s*S\d", r"Figure\s+S\d", r"Table\s+S\d",
    r"SN\s+3\.", r"SN\s+4\.", r"SN\s+5\.",
    r"Additional file 1:", r"Additional file",
    r"Supplementary Note 3\.11",
    r"\([A-E]\)", r"Fig\.\s*\d+[A-E]\b",
    r"Genome Biology", r"NAR-compliant", r"\bNAR\b(?!\s+20\d\d)",
]
ok = True
for pat in residual_pats:
    hits = re.findall(pat, text)
    if hits:
        ok = False
        print(f"RESIDUAL [{pat}]: {len(hits)} -> {hits[:5]}")

n_note = len(re.findall(r"Supplementary Note\s+\d+", text))
notes = sorted(set(re.findall(r"Supplementary Note\s+\d+", text)))
print(f"'Supplementary Note N' occurrences: {n_note} (expected 1); distinct: {notes}")

sec311 = "Section 3.11 of the Supplementary Information (Parameter Justification)"
n311 = text.count(sec311)
print(f"'{sec311}': {n311} (expected 1)")
ok = ok and n311 == 1

figs = sorted(set(re.findall(r"Supplementary Fig\.\s*\d+", text)))
print(f"distinct 'Supplementary Fig. N': {figs}")

supp_info = len(re.findall(r"Supplementary Information", text))
print(f"'Supplementary Information' occurrences: {supp_info}")

# 文件路径中的 Supplementary_Figure_S13.pdf 属真实产物路径，允许保留
path_hits = re.findall(r"Supplementary_Figure_S13\.pdf", text)
print(f"artifact path Supplementary_Figure_S13.pdf kept: {len(path_hits)}")

print("VERIFY", "PASS" if ok and n_note == 1 else "FAIL")
sys.exit(0 if ok and n_note == 1 else 1)
