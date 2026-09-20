# -*- coding: utf-8 -*-
"""verify_nc_ms.py — 用 python-docx 提取 CKI_Manuscript_NC.docx 全文并做 NC 合规自检，
随后写 _tmp_fa_review/nc_ms_conversion_report_2026-09-17.md。"""
import json
import re
from pathlib import Path
from docx import Document

ROOT = Path(r"C:\Users\KnightZ\Desktop\细胞受选择")
DOCX = ROOT / "results" / "CKI_Manuscript_NC.docx"
COUNTS_P = ROOT / "_tmp_fa_review" / "nc_ms_build_counts.json"
REPORT_P = ROOT / "_tmp_fa_review" / "nc_ms_conversion_report_2026-09-17.md"

doc = Document(str(DOCX))
paras = [(p.style.name, p.text) for p in doc.paragraphs]
full = "\n".join(t for _, t in paras)

checks = []
def check(name, ok, detail=""):
    checks.append((name, bool(ok), detail))
    print(("PASS" if ok else "FAIL"), name, ("| " + detail if detail else ""))

# ---- 1. heading order (level 0 Title + Heading 1) ----
h1 = [t for s, t in paras if s in ("Title", "Heading 1")]
expected_h1 = [
    "CKI is a Ka/Ks-inspired index quantifying functional divergence in single-cell genomics",
    "Abstract", "Introduction", "Results", "Discussion", "Methods",
    "Data availability", "Code availability", "References",
    "Acknowledgements", "Author contributions", "Competing interests",
    "Figure legends",
]
check("一级标题顺序 = NC 固定顺序", h1 == expected_h1, " | ".join(h1))

# ---- 2. banned sections / strings ----
for bad in ["Keywords:", "Conclusions", "List of abbreviations", "Declarations",
            "Additional file 1", "Fig. S", "Table S", "Additional file 1:",
            "heading('Background'", "Statistical reporting", "Panel A", "Panel B",
            "Fig. 2A", "Fig. 2B", "Fig. 2C", "(A)", "(B)", "(C)", "(D)", "(E)",
            "Background", "_refs_nar"]:
    if bad == "Conclusions":
        n = sum(1 for s, t in paras if s.startswith("Heading") and t == "Conclusions")
        check("无 Conclusions 标题", n == 0, f"headings named Conclusions={n}")
    elif bad == "Background":
        n = sum(1 for s, t in paras if s.startswith("Heading") and t == "Background")
        check("无 Background 标题", n == 0, f"headings named Background={n}")
    elif bad == "Declarations":
        n = sum(1 for s, t in paras if s.startswith("Heading") and t == "Declarations")
        check("无 Declarations 标题", n == 0)
    elif bad == "_refs_nar":
        continue
    else:
        check(f"全文无 '{bad}'", bad not in full, f"count={full.count(bad)}")

check("无 'Note 3.' / 'Note 4.' / 'Note 5.' 残留",
      not re.search(r"Note [345]\.\d", full), f"matches={re.findall(r'Note [345]\\.\\d', full)}")

# ---- 3. abstract word count ----
abs_i = next(i for i, (s, t) in enumerate(paras) if t == "Abstract")
abs_words = len(paras[abs_i + 1][1].split())
check("Abstract <=150 词", abs_words <= 150, f"{abs_words} words")

# ---- 4. title ----
title = paras[0][1]
check("标题无冒号", ":" not in title, title)
check("标题 <=15 词", len(title.split()) <= 15, f"{len(title.split())} words")

# ---- 5. Additional file count ----
n_af = full.count("Additional file")
check("'Additional file' 出现 <=3", n_af <= 3, f"count={n_af} (expect 2: Methods AF2 + SI paragraph)")
n_af2 = full.count("Additional file 2: Reproducibility Guide")
check("AF2 统一命名", n_af2 == 2, f"count={n_af2}")

# ---- 6. Supplementary naming ----
for n in range(1, 16):
    m = len(re.findall(r"Supplementary Note " + str(n) + r"(?!\d)", full))
    check(f"Supplementary Note {n} 出现", m >= 1, f"count={m}")
n_snote = len(re.findall(r"Supplementary Note \d+", full))
check("Supplementary Note 引用总数 = 38", n_snote == 38, f"count={n_snote}")
n_sfig = len(re.findall(r"Supplementary Fig\. \d+", full))
check("Supplementary Fig. 引用 = 31 (18 正文 + 13 图注)", n_sfig == 31, f"count={n_sfig}")
n_stab = len(re.findall(r"Supplementary Table \d+", full))
check("Supplementary Table 引用 = 4", n_stab == 4, f"count={n_stab}")

# ---- 7. bracket citations gone, CI brackets kept ----
def is_cite(content):
    try:
        nums = [int(x) for x in re.split(r"[,\-]", content)]
    except ValueError:
        return False
    return all(1 <= x <= 56 for x in nums)
bracket_groups = re.findall(r"\[(\d+(?:\s*[,\-]\s*\d+)*)\]", full)
cite_like = [g for g in bracket_groups if is_cite(g)]
check("[n] 方括号引用组 = 0", len(cite_like) == 0, f"citation-like={cite_like[:5]}")
check("CI 方括号保留 (>0)", len(bracket_groups) > 0, f"int-bracket groups={len(bracket_groups)}")

# ---- 8. superscript runs ----
n_sup = 0
for p in doc.paragraphs:
    for r in p.runs:
        if r.font.superscript:
            n_sup += 1
check("上标 run 总数 = 77 (8 标题页 + 69 引用)", n_sup == 77, f"count={n_sup}")

# ---- 9. Data/Code availability: no citations ----
da_i = next(i for i, (s, t) in enumerate(paras) if t == "Data availability")
ref_i = next(i for i, (s, t) in enumerate(paras) if t == "References")
avail_paras = doc.paragraphs[da_i:ref_i]
avail_text = "\n".join(p.text for p in avail_paras)
avail_cites = [g for g in re.findall(r"\[(\d+(?:\s*[,\-]\s*\d+)*)\]", avail_text) if is_cite(g)]
avail_sup = sum(1 for p in avail_paras for r in p.runs if r.font.superscript)
check("Data/Code availability 内无 [n] 引用", len(avail_cites) == 0, f"{avail_cites}")
check("Data/Code availability 内无上标 run", avail_sup == 0, f"{avail_sup}")
check("Data availability 含 GSE109774", "GSE109774" in avail_paras[1].text)
check("Code availability 含 Zenodo DOI", "10.5281/zenodo.20405458" in avail_paras[3].text)

# ---- 10. References ----
ack_i = next(i for i, (s, t) in enumerate(paras) if t == "Acknowledgements")
ref_paras = [p for p in doc.paragraphs[ref_i + 1:ack_i] if p.text.strip()]
check("References 条目 = 56", len(ref_paras) == 56, f"count={len(ref_paras)}")
n_year = sum(1 for p in ref_paras if re.search(r"\(\d{4}\)\.$", p.text.strip()))
check("条目以 (年份). 结尾 = 55（#54 书籍条目为例外）", n_year == 55, f"count={n_year}")
r2 = ref_paras[1]
check("条目2 = Korsunsky Nature 格式",
      r2.text == "2. Korsunsky, I. et al. Fast, sensitive and accurate integration of single-cell data with Harmony. Nat. Methods 16, 1289\u20131296 (2019).",
      r2.text)
ital = [r.text for r in r2.runs if r.italic]
bold = [r.text for r in r2.runs if r.bold]
check("条目2 期刊斜体", ital == ["Nat. Methods"], f"italic={ital}")
check("条目2 卷号加粗", bold == ["16,"], f"bold={bold}")
r54 = ref_paras[53]
check("条目54 = Efron & Tibshirani 书籍格式",
      r54.text == "54. Efron, B. & Tibshirani, R. J. An Introduction to the Bootstrap (Chapman and Hall/CRC, 1994).",
      r54.text)
et_al6 = ["4.", "20.", "27.", "29.", "37.", "56."]
trunc = all(ref_paras[int(n[:-1]) - 1].text.split(". ", 1)[1].strip().endswith("et al.") or " et al. " in ref_paras[int(n[:-1]) - 1].text for n in et_al6)
check("6 条六作者条目已截断为第 1 作者 + et al. (#4/#20/#27/#29/#37/#56)", trunc)
n_vancouver = sum(1 for p in ref_paras if re.search(r"\d{4};\d+:", p.text))
check("无 GB/Vancouver 式 '年;卷:页' 残留", n_vancouver == 0, f"count={n_vancouver}")

# ---- 11. heading levels ----
n_h3 = sum(1 for s, t in paras if s == "Heading 3")
check("无 Heading 3 (Results L3 已删)", n_h3 == 0, f"count={n_h3}")
disc_i = next(i for i, (s, t) in enumerate(paras) if t == "Discussion")
meth_i = next(i for i, (s, t) in enumerate(paras) if t == "Methods")
n_h2_disc = sum(1 for s, t in paras[disc_i:meth_i] if s == "Heading 2")
check("Discussion 内无小节标题", n_h2_disc == 0, f"count={n_h2_disc}")
n_stats = sum(1 for s, t in paras if s == "Heading 2" and t == "Statistics and reproducibility")
check("Methods 含 'Statistics and reproducibility'", n_stats == 1)
n_h2 = [(t) for s, t in paras if s == "Heading 2"]
check("L2 小节总数 = 25 (Results 11 + Methods 14)", len(n_h2) == 25, f"count={len(n_h2)}")

# ---- 12. declarations content ----
ack_p = paras[ack_i + 1][1]
check("Funding 并入 Acknowledgements 段首", ack_p.startswith("This work was supported by the National Natural Science Foundation of China (NSFC) under grant number 32370682."), ack_p[:80])
ci_i = next(i for i, (s, t) in enumerate(paras) if t == "Competing interests")
check("Competing interests 声明", paras[ci_i + 1][1] == "The authors declare no competing interests.", paras[ci_i + 1][1])
n_si = full.count("Supplementary Information is available for this paper")
check("Supplementary Information 描述段", n_si == 1)

# ---- 13. legends ----
fl_i = next(i for i, (s, t) in enumerate(paras) if t == "Figure legends")
legend_paras = [t for s, t in paras[fl_i + 1:] if t.strip()]
main_figs = [t for t in legend_paras if re.match(r"Figure \d\.", t)]
supp_figs = [t for t in legend_paras if re.match(r"Supplementary Fig\. \d+\.", t)]
check("主图图注 = 6", len(main_figs) == 6, f"count={len(main_figs)}")
check("附图图注 = 13", len(supp_figs) == 13, f"count={len(supp_figs)}")
check("附图图注均为 'Supplementary Fig. n.' 命名",
      [re.match(r"Supplementary Fig\. (\d+)\.", t).group(1) for t in supp_figs] == [str(i) for i in range(1, 14)])

# ---- 14. spot checks ----
check("正文含 'Supplementary Fig. 1'", "Supplementary Fig. 1" in full)
check("正文含 'Fig. 2b, c'", "Fig. 2b, c" in full)
check("正文含 'panel a'", "panel a" in full)
check("正文引用 '(X)' 保留", "(X)" in full)

# ---- 15. round-2 fixes ----
intro_i = next(i for i, (s, t) in enumerate(paras) if t == "Introduction")
res_i = next(i for i, (s, t) in enumerate(paras) if t == "Results")
intro_paras = [t for s, t in paras[intro_i + 1:res_i] if t.strip()]
check("Introduction 末段以 'Here, we show' 开头",
      intro_paras[-1].startswith("Here, we show"), intro_paras[-1][:90])
long_h2 = [t for s, t in paras if s == "Heading 2" and len(t) > 60]
check("全部 L2 标题 <=60 字符", len(long_h2) == 0, f"over={long_h2}")
for h in ["Ground-truth simulation: specificity versus sensitivity",
          "Fixed-panel ablation: robust rankings, scheme-specific \u03c9",
          "Cancer analysis: apparent tumor homogeneity (exploratory)",
          "Brain regional analysis reveals divergence gradients",
          "Anomalously similar pairs: a hypothesis-generating screen"]:
    check(f"新短标题到位: {h[:45]}", any(s == "Heading 2" and t == h for s, t in paras))
check("无悬空 '(Results; Limitations)'", "(Results; Limitations)" not in full)
check("无悬空 '; Limitations)'", "; Limitations)" not in full)
check("'Limitations' 大写残留 = 1（仅 'Limitations replicated too' 概念句）",
      full.count("Limitations") == 1, f"count={full.count('Limitations')}")

# ================= report =================
counts_data = json.loads(COUNTS_P.read_text(encoding="utf-8"))
old_refs = counts_data["old_refs"]
new_refs = counts_data["new_refs"]
def demark(s):
    return (s.replace("\u00abi\u00bb", "*").replace("\u00ab/i\u00bb", "*")
             .replace("\u00abb\u00bb", "**").replace("\u00ab/b\u00bb", "**"))

n_pass = sum(1 for _, ok, _ in checks if ok)
n_fail = sum(1 for _, ok, _ in checks if not ok)

lines = []
lines.append("# CKI 文稿 GB→NC 生成器改造报告（nc-ms）")
lines.append("")
lines.append("- 日期：2026-09-17")
lines.append("- 实施脚本：`_tmp_fa_review/build_nc_ms.py`（含全部锚点断言）")
lines.append("- 产出生成器：`generate_manuscript_nc.py`（py_compile 通过，运行成功）")
lines.append("- 产出文稿：`results/CKI_Manuscript_NC.docx`")
lines.append("- 验证脚本：`_tmp_fa_review/verify_nc_ms.py`（python-docx 全文提取自检）")
lines.append(f"- **自检结果：{n_pass} 项通过，{n_fail} 项失败**")
lines.append("")
lines.append("## 一、各项改造替换计数")
lines.append("")
lines.append("| 改造项 | 计数 |")
lines.append("|---|---|")
c = counts_data["counts"]
rows = [
    ("结构：Background→Introduction 标题", c["heading Background->Introduction"]),
    ("结构：删除 Conclusions 标题行（段落并入 Discussion 末）", c["delete Conclusions heading"]),
    ("结构：删除 Keywords 块", c["delete Keywords block"]),
    ("结构：删除 List of abbreviations 块", c["delete List of abbreviations"]),
    ("结构：删除 Declarations/Ethics/Consent 块", c["delete Declarations/Ethics/Consent"]),
    ("结构：删除 Additional files 块（改写为 SI 描述段）", c["delete Additional files block"]),
    ("结构：删除附图图例标题（附图图注并入 Figure legends 节）", c["delete supp legends heading"]),
    ("结构：删除旧 References 循环（移至 Code availability 之后）", c["delete old References loop (moved)"]),
    ("标题页：去冒号新标题（11 词）", c["title (no colon, <=15 words)"]),
    ("Abstract：151→149 词", c["abstract trim"]),
    ("Results：删除 L3 小节标题", 3),
    ("Discussion：删除 L2 小节标题", 3),
    ("Methods：Statistical reporting→Statistics and reproducibility（含正文交叉引用 1 处）", 2),
    ("正文残留 'Background' 引用→Introduction", c["body (Background)->(Introduction)"]),
    ("命名：'Additional file 1: Fig. Sx'→'Supplementary Fig. x'", c["AF1: Fig. Sx -> Supplementary Fig."]),
    ("命名：裸 'Fig. Sx'→'Supplementary Fig. x'", c["bare Fig. Sx -> Supplementary Fig."]),
    ("命名：图注 'Additional file 1: Figure Sx.'→'Supplementary Fig. x.'", c["AF1: Figure Sx. -> Supplementary Fig."]),
    ("命名：'Additional file 1: Table Sx'→'Supplementary Table x'", c["AF1: Table Sx -> Supplementary Table"]),
    ("命名：'Additional file 1: Note x.y'→'Supplementary Note 新号'（前缀式）", sum(v for k, v in c.items() if k.startswith("AF1: Note"))),
    ("命名：裸 'Note x.y'→'Supplementary Note 新号'", sum(v for k, v in c.items() if k.startswith("bare Note"))),
    ("命名：'in Additional file 1.'→'in the Supplementary Information.'", c["bare 'in Additional file 1.' -> Supplementary Information"]),
    ("命名：Additional file 2→'Additional file 2: Reproducibility Guide'", c["AF2 reference renamed"]),
    ("面板标签：(A)-(E)→(a)-(e)", c["(A)->(a)"] + c["(B)->(b)"] + c["(C)->(c)"] + c["(D)->(d)"] + c["(E)->(e)"]),
    ("面板标签：Fig. 2A/2B/2C→小写（含 'Fig. 2B, C'→'Fig. 2b, c'）", 3),
    ("面板标签：Panel A/B→panel a/b", 2),
    ("参考文献：56 条 Nature 化（含 6 条六作者截断、书籍/团体作者套新框架）", 56),
    ("正文引用：72 组 [n] → 上标 69 组（3 组随 Data/Code availability 去引用删除）", 69),
    ("p() 函数：引用上标化改造（CI 方括号守卫）", c["p() citation superscript"]),
    ("ref_p_nar→ref_p_nc：期刊斜体/卷号加粗 run 拆分", c["ref_p_nar -> ref_p_nc"]),
    ("保存路径→results/CKI_Manuscript_NC.docx", c["save path"]),
    ("第二轮：Introduction 末段 'Here, we show' 改写", c["intro final paragraph 'Here, we show'"]),
    ("第二轮：Results 超长 L2 标题缩短 ≤60 字符", 5),
    ("第二轮：悬空交叉引用改写（指向已删小节）", 3),
]
for name, n in rows:
    lines.append(f"| {name} | {n} |")
lines.append("")
lines.append("Supplementary Note 旧→新映射（引用次数）：3.12→1(5), 3.21→2(4), 3.5→3(2), 3.20→4(1), 3.22→5(2), 3.15→6(2), 3.13→7(2), 5.2→8(2), 3.16→9(6), 3.17→10(2), 3.14→11(2), 4.6→12(3), 5.1→13(3), 3.23→14(1), 3.6→15(1)，合计 38。")
lines.append("")
lines.append("## 二、DOCX 自检明细")
lines.append("")
lines.append("| 检查项 | 结果 | 细节 |")
lines.append("|---|---|---|")
for name, ok, detail in checks:
    lines.append(f"| {name} | {'PASS' if ok else 'FAIL'} | {detail.replace('|', '/')[:120]} |")
lines.append("")
lines.append("## 三、参考文献 56 条新旧对照表")
lines.append("")
lines.append("| # | GB (Vancouver) | NC (Nature) |")
lines.append("|---|---|---|")
for i, (o, n_) in enumerate(zip(old_refs, new_refs), 1):
    lines.append(f"| {i} | {o} | {demark(n_)} |")
lines.append("")
lines.append("## 四、第二轮补修（2026-09-17，team-lead 指示）")
lines.append("")
lines.append("### 4.1 Introduction 末段改写（NC 惯例 'Here, we show' 开头）")
lines.append("")
lines.append("- 旧：`We evaluated CKI across four scales. First, we calibrated CKI on Tabula Muris mouse data …`")
lines.append("- 新：`Here, we show that CKI provides a baseline-normalized index of cell-state divergence across four scales. First, we calibrated CKI on Tabula Muris mouse data …`（其余文本不动）")
lines.append("")
lines.append("### 4.2 Results 超长 L2 标题缩短（≤60 字符）新旧对照")
lines.append("")
lines.append("| 旧标题（字符数） | 新标题（字符数） |")
lines.append("|---|---|")
h_pairs = [
    ("Ground-truth simulation dissociates specificity from sensitivity", "Ground-truth simulation: specificity versus sensitivity"),
    ("Fixed gene-panel ablation: rankings robust, absolute \u03c9 scheme-specific", "Fixed-panel ablation: robust rankings, scheme-specific \u03c9"),
    ("Cancer analysis suggests apparent tumor homogeneity (exploratory)", "Cancer analysis: apparent tumor homogeneity (exploratory)"),
    ("Brain regional analysis reveals cell-type divergence gradients", "Brain regional analysis reveals divergence gradients"),
    ("Anomalously similar cell-type/region pairs: a hypothesis-generating screen", "Anomalously similar pairs: a hypothesis-generating screen"),
]
for old_h, new_h in h_pairs:
    lines.append(f"| {old_h}（{len(old_h)}） | {new_h}（{len(new_h)}） |")
lines.append("")
lines.append("### 4.3 悬空交叉引用改写（指向已删 Discussion 小节）")
lines.append("")
lines.append("| 位置 | 旧 | 新 |")
lines.append("|---|---|---|")
lines.append("| Introduction L4 段 | `(Results; Limitations)` | `(Results; limitations are addressed in the Discussion)` |")
lines.append("| Results（benchmarking 段） | `(operating window ~50–200 cells; Limitations)` | `(operating window ~50–200 cells; limitations are addressed in the Discussion)` |")
lines.append("| Discussion（confounder 段） | `RNA quality; Limitations)` | `RNA quality; limitations discussed below)` |")
lines.append("")
lines.append("注：Results 模拟段 'Limitations replicated too: …' 为句首概念用法（非交叉引用），保留未动。")
lines.append("")
lines.append("## 五、遗留事项（提请 team-lead 决策）")
lines.append("")
lines.append("1. 正文 Intro+Results+Discussion 12,151 词、Methods 4,437 词，超 NC 指南值 5,000/3,000（nc-gap §1.4）——team-lead 已决定按现状投。")
lines.append("2. 图 PDF 内面板标签大小写未审计（figure1-6.pdf，nc-gap §六）——本轮不改。")
lines.append("3. Nature Portfolio Reporting Summary 表单——team-lead 另行处理。")
lines.append("4. 断言体系（verify_v40 式）中引用已删标题文本的断言需由 build 工人更新。")

REPORT_P.write_text("\n".join(lines), encoding="utf-8")
print(f"\nReport: {REPORT_P}")
print(f"CHECKS: {n_pass} pass, {n_fail} fail")
if n_fail:
    raise SystemExit(1)
