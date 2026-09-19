"""
Generate Nature Communications Cover Letter (v49 repositioning)
CKI Project — Nature Communications submission (Article)

v49 changes (d-tcga, 2026-09-18):
  * Value proposition rewritten: from "a better metric" to the first
    principled separation of functional divergence from background drift
    (Ka/Ks analogy), with real-data misreporting-control evidence and a
    discovery-level pan-cancer application.
  * Reviewer-concern paragraph replaced with a new-evidence framing ("Two
    properties of this work are worth stating explicitly..."), per blind
    review R4 P0-3/P1-1: no rebuttal voice, no "5th of 5".
  * Misreporting claim qualified to the continuous divergence metrics
    (marker Jaccard is lower still on the FPR statistic), per R1 P1-1.
  * Preserved: NC scope, six suggested reviewers, Zenodo DOI, one-page
    format (Arial 11 pt, 1" margins, ~520 words).
Numbers sourced from verified audits only:
  nc49_pilot_kang_2026-09-18.md (Kang 30 pairs: omega 0/30 vs raw JS 36.7%,
  cosine 23.3%), nc49_brain_ladder_2026-09-18.md (T1 2,161 pairs: omega FPR
  28.6% lowest among continuous metrics, below raw JS 10/10 classes; ladder
  1.04->1.76->1.80 vs raw JS 1.07->3.32->2.98),
  nc49_tcga_main_2026-09-18.md (3,567 samples; NN/TT
  1.10-2.46, 4/5 CIs exclude 1; k_n 1.3-3.3x; LUAD KRAS retained after
  purity + smoking adjustment, EGFR association dissolved under purity
  adjustment per nc49_purity/nc49_smoking audits),
  manuscript Table 1 (classification AUC 0.680) and simulation
  results (FPR 0.00 vs 0.55-0.58; AUC 0.80).
"""
from pathlib import Path
from docx.shared import Pt, Inches

PROJECT_ROOT = Path(__file__).parent
OUTPUT_DIR = PROJECT_ROOT / "results"
OUTPUT_DIR.mkdir(exist_ok=True)


def add_para(text, doc, space_after=4, align=None, bold=False, size=None):
    """Add a paragraph with controlled spacing — no empty-paragraph spacers."""
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 1.0
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(0)
    run = para.add_run(text)
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if align is not None:
        para.alignment = align
    return para


def run():
    from docx import Document

    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Salutation (letter begins directly at "Dear Editors") ──
    add_para("Dear Editors,", doc, space_after=6)

    # ── Body: opening + NC scope + repositioned value proposition ──
    add_para(
        "On behalf of my co-author, Dr. Xianming Wu (Chinese Institute for "
        "Brain Research, Beijing), I submit our manuscript, "
        "\u201cCKI is a Ka/Ks-inspired index quantifying functional divergence "
        "in single-cell genomics\u201d, for consideration as an Article in "
        "Nature Communications. CKI (Cell-type Ka/Ks-inspired Index) is, to "
        "our knowledge, the first framework to make a principled separation "
        "of functional divergence from background drift in transcriptomic "
        "comparisons: adapting the "
        "Ka/Ks ratio\u2019s logic, it decomposes divergence into a "
        "housekeeping baseline k_n and an identity-gene rate k_f "
        "(\u03c9 = k_f/k_n). The work fits Nature Communications\u2019 scope "
        "of computational methods validated on large-scale real data.",
        doc,
    )

    # ── Body: core argument — three pillars ──
    add_para(
        "The manuscript rests on three pillars. First, in ground-truth "
        "simulation (1,750 replicates), \u03c9 alone did not false-trigger on "
        "neutral housekeeping drift (false-positive rate 0.00 versus "
        "0.55\u20130.58 for JS and cosine) and ranked first for "
        "functional-versus-neutral discrimination (AUC = 0.80). Second, in "
        "real-data drift calibration, \u03c9 raised no false reports on any of the 30 "
        "cross-lane technical-replicate pairs (Kang IFN-\u03b2 PBMC data; "
        "raw JS 36.7%, cosine 23.3%) and, on 2,161 brain technical-drift "
        "pairs, had the lowest false-report rate among the continuous "
        "divergence metrics (28.6% versus 35.7\u201345.2% for the other "
        "five; below raw JS in all ten cell classes), with the shallowest "
        "drift-ladder gradient (1.04 \u2192 1.76 \u2192 1.80 versus "
        "1.07 \u2192 3.32 \u2192 2.98 for raw JS). Third, across 3,567 TCGA "
        "samples in five cancer types, tumor specimens were consistently less "
        "divergent than adjacent non-tumor tissue (NN/TT \u03c9 ratio "
        "1.10\u20132.46, bootstrap CIs excluding 1 in four of five)\u2014"
        "reflecting a 1.3\u20133.3-fold elevated housekeeping baseline, not "
        "reduced functional divergence; and in lung adenocarcinoma the k_f/k_n "
        "decomposition separates KRAS-mutant tumors\u2014retaining both a "
        "functional (k_f) and a baseline (k_n) component after purity and "
        "smoking adjustment\u2014from EGFR-mutant tumors, whose apparent "
        "association dissolved under purity adjustment.",
        doc,
    )

    # ── Body: two properties stated explicitly (new-evidence framing) ──
    add_para(
        "Two properties of this work are worth stating explicitly. CKI is "
        "a specificity-first index: its cell-type classification AUC "
        "(0.680) is modest by design, because it down-weights the "
        "global-identity signal that classifiers exploit; on the benchmark "
        "matched to its question domain\u2014false divergence calls on "
        "real technical and donor drift\u2014it misreports least among "
        "the continuous divergence metrics while retaining sensitivity to "
        "biology (Fig. 4). And the pan-cancer divergence map, with the "
        "lung adenocarcinoma driver-class decomposition, grounds the "
        "method in cancer biology of direct interest to a broad "
        "readership (Fig. 5).",
        doc,
    )

    # ── Body: declarations, reproducibility, funding ──
    add_para(
        "Both authors declare no competing interests. The work is original, "
        "not under consideration elsewhere, and not previously submitted to "
        "Nature Communications. AI tools were used for debugging, code "
        "review, and language editing; all AI-assisted content was reviewed "
        "and revised by the authors, who take full responsibility. This work "
        "was supported by the National Natural Science Foundation of China "
        "(grant 32370682). The CKI Python package (v0.5.0, MIT License) and "
        "all analysis code are available at "
        "https://github.com/zhanglknt/CKI-cell-type-identification (Zenodo "
        "DOI 10.5281/zenodo.22735744).",
        doc,
    )

    # ── Body: suggested reviewers ──
    add_para(
        "We suggest the following reviewers, none with recent collaborations "
        "or conflicts of interest with the authors: "
        "Prof. Fabian Theis, Helmholtz Munich (fabian.theis@helmholtz-munich.de); "
        "Prof. Joshua Welch, University of Michigan (welchjd@umich.edu); "
        "Prof. Sten Linnarsson, Karolinska Institutet (sten.linnarsson@ki.se); "
        "Prof. Patrik St\u00e5hl, KTH / SciLifeLab (patrik.stahl@scilifelab.se); "
        "Dr. Alejandro A. Sch\u00e4ffer, National Cancer Institute, NIH "
        "(schaffer@ncbi.nlm.nih.gov); and Prof. Zemin Zhang, Peking University "
        "(zeminzhang@pku.edu.cn).",
        doc,
    )

    add_para(
        "Thank you for considering our work.",
        doc,
    )

    # ── Closing ──
    add_para("Sincerely,", doc, space_after=12)
    add_para("Li Zhang (Corresponding Author)", doc, space_after=0)
    add_para("Xianming Wu (First Author)", doc, space_after=0)

    # Save
    out = str(OUTPUT_DIR / "CKI_NC_Cover_Letter.docx")
    doc.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    run()
